"""
Build ClaimsKPI_Report.xlsx from the KPI views in SQL Server.

One sheet per view (a snapshot of the view's rows) plus a Summary sheet whose
chart tables are SUMIFS formulas over those sheets. Because the formulas use
whole-column references, the same Summary works when the view sheets are
replaced by live Excel queries (see docs/excel_report_setup.md), and it
recalculates when the file is imported into Google Sheets.

Usage:  py -3.9 scripts/build_report.py
Needs:  pip install pyodbc openpyxl
"""
import datetime as dt
import decimal
import pathlib

import pyodbc
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = pathlib.Path(__file__).resolve().parent.parent / "ClaimsKPI_Report.xlsx"
CONN = ("DRIVER={ODBC Driver 18 for SQL Server};SERVER=localhost\\SQLEXPRESS;"
        "DATABASE=ClaimsKPI;Trusted_Connection=yes;TrustServerCertificate=yes")

# sheet name -> query (sheet names match the view names)
VIEWS = {
    "v_kpi_monthly":             "SELECT * FROM dbo.v_kpi_monthly ORDER BY month_start, claim_type",
    "v_kpi_weekly":              "SELECT * FROM dbo.v_kpi_weekly ORDER BY week_start, claim_type",
    "v_kpi_annual_pmpm":         "SELECT * FROM dbo.v_kpi_annual_pmpm ORDER BY yr, claim_type",
    "v_kpi_inpatient_quarterly": "SELECT * FROM dbo.v_kpi_inpatient_quarterly ORDER BY yr, qtr",
    "v_inpatient_stays":         "SELECT * FROM dbo.v_inpatient_stays ORDER BY member_id, admit_dt",
    "v_paid_by_diagnosis":       "SELECT * FROM dbo.v_paid_by_diagnosis ORDER BY total_paid DESC",
}

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
MONEY, PCT, DATE = '"$"#,##0', "0.0%", "yyyy-mm-dd"


def style_header(ws, row, first_col, last_col):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL


def write_view(wb, cur, name, sql):
    ws = wb.create_sheet(name)
    cur.execute(sql)
    cols = [d[0] for d in cur.description]
    ws.append(cols)
    style_header(ws, 1, 1, len(cols))
    for row in cur.fetchall():
        ws.append([float(v) if isinstance(v, decimal.Decimal) else v for v in row])
    for i, col in enumerate(cols, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(12, len(col) + 2)
        fmt = (DATE if col.endswith(("_start", "_dt")) else
               MONEY if col in ("total_paid", "avg_paid_per_claim", "paid_amount", "pmpm") else
               PCT if col == "readmit_30_rate" else None)
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    return ws.max_row - 1


def build_summary(ws):
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 30
    for col in "CDEFG":
        ws.column_dimensions[col].width = 14

    # B1 is reserved for RefreshAndExport's "Last refreshed" stamp
    ws["B1"] = f"Last refreshed: {dt.datetime.now():%Y-%m-%d %H:%M}"
    ws["B1"].font = Font(italic=True, color="666666")
    ws["B2"] = "Claims KPI Summary: CMS DE-SynPUF Sample 1, 2008-2010"
    ws["B2"].font = Font(bold=True, size=14)

    # Headline KPIs (2008-2010)
    ws["B4"] = "Headline KPIs (2008-2010)"
    ws["B4"].font = Font(bold=True)
    headline = [
        ("Total paid (IP + OP)", "=SUM(v_kpi_monthly!E:E)", MONEY),
        ("PMPM (IP + OP)", '=SUM(v_kpi_annual_pmpm!C:C)/SUMIFS(v_kpi_annual_pmpm!D:D,v_kpi_annual_pmpm!B:B,"IP")', '"$"#,##0.00'),
        ("Inpatient admissions", "=SUM(v_kpi_inpatient_quarterly!C:C)", "#,##0"),
        ("Average length of stay (days)", '=AVERAGEIFS(v_inpatient_stays!F:F,v_inpatient_stays!D:D,">="&DATE(2008,1,1))', "0.00"),
        ("30-day readmission rate", '=AVERAGEIFS(v_inpatient_stays!G:G,v_inpatient_stays!D:D,">="&DATE(2008,1,1))', PCT),
    ]
    for i, (label, formula, fmt) in enumerate(headline, start=5):
        ws.cell(row=i, column=2, value=label)
        ws.cell(row=i, column=3, value=formula).number_format = fmt

    # Table 1: monthly paid (rows 12-48)
    t1 = 12
    ws.cell(row=t1 - 1, column=2, value="Monthly paid amount").font = Font(bold=True)
    for j, h in enumerate(["Month", "Inpatient", "Outpatient", "Total"], start=2):
        ws.cell(row=t1, column=j, value=h)
    style_header(ws, t1, 2, 5)
    for k in range(36):
        r = t1 + 1 + k
        ws.cell(row=r, column=2, value=dt.date(2008 + k // 12, k % 12 + 1, 1)).number_format = "mmm yyyy"
        for col, ctype in ((3, "IP"), (4, "OP")):
            ws.cell(row=r, column=col,
                    value=f'=SUMIFS(v_kpi_monthly!$E:$E,v_kpi_monthly!$A:$A,$B{r},v_kpi_monthly!$B:$B,"{ctype}")'
                    ).number_format = MONEY
        ws.cell(row=r, column=5, value=f"=C{r}+D{r}").number_format = MONEY
    t1_end = t1 + 36

    # Table 2: PMPM by year
    t2 = t1_end + 3
    ws.cell(row=t2 - 1, column=2, value="PMPM by year").font = Font(bold=True)
    for j, h in enumerate(["Year", "Inpatient", "Outpatient", "Total"], start=2):
        ws.cell(row=t2, column=j, value=h)
    style_header(ws, t2, 2, 5)
    for k, yr in enumerate((2008, 2009, 2010)):
        r = t2 + 1 + k
        ws.cell(row=r, column=2, value=str(yr))
        for col, ctype in ((3, "IP"), (4, "OP")):
            ws.cell(row=r, column=col,
                    value=f'=SUMIFS(v_kpi_annual_pmpm!$E:$E,v_kpi_annual_pmpm!$A:$A,{yr},v_kpi_annual_pmpm!$B:$B,"{ctype}")'
                    ).number_format = '"$"#,##0.00'
        ws.cell(row=r, column=5, value=f"=C{r}+D{r}").number_format = '"$"#,##0.00'
    t2_end = t2 + 3

    # Table 3: readmission rate by quarter
    t3 = t2_end + 3
    ws.cell(row=t3 - 1, column=2, value="30-day readmission rate by quarter").font = Font(bold=True)
    for j, h in enumerate(["Quarter", "Readmit rate", "Admissions", "Year", "Qtr"], start=2):
        ws.cell(row=t3, column=j, value=h)
    style_header(ws, t3, 2, 6)
    for k in range(12):
        r, yr, q = t3 + 1 + k, 2008 + k // 4, k % 4 + 1
        ws.cell(row=r, column=2, value=f"{yr} Q{q}")
        ws.cell(row=r, column=5, value=yr)
        ws.cell(row=r, column=6, value=q)
        crit = f"v_kpi_inpatient_quarterly!$A:$A,$E{r},v_kpi_inpatient_quarterly!$B:$B,$F{r}"
        ws.cell(row=r, column=3, value=f"=SUMIFS(v_kpi_inpatient_quarterly!$F:$F,{crit})").number_format = PCT
        ws.cell(row=r, column=4, value=f"=SUMIFS(v_kpi_inpatient_quarterly!$C:$C,{crit})").number_format = "#,##0"
    t3_end = t3 + 12

    # Charts
    c1 = LineChart()
    c1.title, c1.y_axis.title, c1.height, c1.width = "Monthly paid amount", "Paid ($)", 8, 18
    c1.y_axis.numFmt = '"$"#,##0'
    c1.add_data(Reference(ws, min_col=3, max_col=5, min_row=t1, max_row=t1_end), titles_from_data=True)
    c1.set_categories(Reference(ws, min_col=2, min_row=t1 + 1, max_row=t1_end))
    ws.add_chart(c1, "H4")

    c2 = BarChart()
    c2.type, c2.grouping, c2.overlap = "col", "stacked", 100
    c2.title, c2.y_axis.title, c2.height, c2.width = "PMPM by year", "PMPM ($)", 8, 18
    c2.y_axis.numFmt = '"$"#,##0'
    c2.add_data(Reference(ws, min_col=3, max_col=4, min_row=t2, max_row=t2_end), titles_from_data=True)
    c2.set_categories(Reference(ws, min_col=2, min_row=t2 + 1, max_row=t2_end))
    ws.add_chart(c2, "H21")

    c3 = LineChart()
    c3.title, c3.y_axis.title, c3.height, c3.width = "30-day readmission rate by quarter", "Rate", 8, 18
    c3.y_axis.numFmt = "0%"
    c3.add_data(Reference(ws, min_col=3, min_row=t3, max_row=t3_end), titles_from_data=True)
    c3.set_categories(Reference(ws, min_col=2, min_row=t3 + 1, max_row=t3_end))
    c3.legend = None
    ws.add_chart(c3, "H38")


def main():
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    with pyodbc.connect(CONN) as cn:
        cur = cn.cursor()
        for name, sql in VIEWS.items():
            print(f"{name:<28} {write_view(wb, cur, name, sql):>7,} rows")
    build_summary(summary)
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
